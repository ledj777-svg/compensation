from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.ingest import read_workbook
from app.sample_workbook import build_sample_workbook
from app.validation import (
    DUPLICATE_RECORD,
    INVALID_RATING,
    MISSING_EMPLOYEE_ID,
    MISSING_GRADE,
    MISSING_SALARY,
)


def build_dirty_workbook() -> bytes:
    workbook = load_workbook(BytesIO(build_sample_workbook()))
    sheet = workbook["Employee_Master"]
    sheet.append([None, "No ID", "Technology", "A1", "Engineer", 10, 4, None, None, 2, 4, "Bengaluru", "M", "P", "Active"])
    sheet.append([9001, "No Salary", "Technology", "A1", "Engineer", None, 4, None, None, 2, 4, "Bengaluru", "M", "P", "Active"])
    sheet.append([9002, "No Grade", "Technology", None, "Engineer", 10, 4, None, None, 2, 4, "Bengaluru", "M", "P", "Active"])
    sheet.append([9003, "Bad Rating", "Technology", "A1", "Engineer", 10, 9, None, None, 2, 4, "Bengaluru", "M", "P", "Active"])
    sheet.append([1001, "Dup Ananya", "Technology", "A1", "Engineer", 12, 4, None, None, 2, 4, "Bengaluru", "M", "P", "Active"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_validation_engine_flags_required_problems():
    employees, _increments, _bands, _corrections, issues, counts = read_workbook(build_dirty_workbook())
    codes = {item.code for item in issues}
    assert MISSING_EMPLOYEE_ID in codes
    assert MISSING_SALARY in codes
    assert MISSING_GRADE in codes
    assert INVALID_RATING in codes
    assert DUPLICATE_RECORD in codes
    ids = {employee.employee_id for employee in employees}
    assert "9001" not in ids
    assert "9002" not in ids
    assert "9003" in ids
    assert counts["Employee_Master"] == 29
    original = next(employee for employee in employees if employee.employee_id == "1001")
    assert original.current_salary == 10
    bad_rating = next(employee for employee in employees if employee.employee_id == "9003")
    assert bad_rating.performance_rating is None


def test_employee_only_workbook_is_accepted():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(
        [
            "Employee ID",
            "Employee Name",
            "Department",
            "Grade",
            "Designation",
            "Current Salary",
            "Performance Rating",
            "Years At Level",
        ]
    )
    sheet.append([1001, "Ananya Rao", "Technology", "A1", "Software Engineer", 10, 4, 2])
    buffer = BytesIO()
    workbook.save(buffer)
    employees, increments, bands, corrections, _issues, counts = read_workbook(buffer.getvalue())
    assert len(employees) == 1
    assert employees[0].employee_id == "1001"
    assert increments == []
    assert bands == []
    assert corrections == []
    assert "Increment_Grid" in counts["missing_sheets"]
    assert "Salary_Band" in counts["missing_sheets"]
    assert "Current Salary" in counts["found_fields"]
    assert counts["complete"] is False


def test_employee_only_uses_predefined_formulas():
    from openpyxl import Workbook

    from app.defaults import apply_defaults
    from app.engine import recommend_employee

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Employee ID",
            "Employee Name",
            "Department",
            "Grade",
            "Designation",
            "Current Salary",
            "Performance Rating",
            "Years At Level",
        ]
    )
    sheet.append([1001, "Ananya Rao", "Technology", "A1", "Software Engineer", 10, 4, 2])
    buffer = BytesIO()
    workbook.save(buffer)
    employees, increments, bands, corrections, _issues, _counts = read_workbook(buffer.getvalue())
    increments, bands, corrections, used = apply_defaults(employees, increments, bands, corrections)
    assert used == ["Increment_Grid", "Salary_Band", "Salary_Correction_Rules"]
    result = recommend_employee(employees[0], bands, increments, corrections)
    assert result.increment_pct == 0.10
    assert result.correction_pct == 0.08
    assert result.recommended_salary == 11.8
