from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import CompensationResult, ValidationIssue

NAVY = "243044"
COPPER = "9A5B28"
PAPER = "FBF7F0"
INK = "1C1612"
LINE = "D7CBB8"
BELOW_MIN = "F4D6D6"
PRIORITY = "F7E6C8"
ABOVE = "E4EAF3"

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Calibri", color=INK, size=11)
MONO_FONT = Font(name="Consolas", color=INK, size=10)
THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


REPORT_COLUMNS = [
    ("Employee ID", "employee_id", None, 14),
    ("Employee Name", "employee_name", None, 22),
    ("Department", "department", None, 16),
    ("Grade", "grade", None, 10),
    ("Designation", "designation", None, 22),
    ("Current Salary", "current_salary", "0.00", 16),
    ("Performance Rating", "performance_rating", "0", 14),
    ("Years At Level", "years_at_level", "0.0", 14),
    ("Minimum Salary", "band_min", "0.00", 16),
    ("Median Salary", "band_median", "0.00", 16),
    ("Maximum Salary", "band_max", "0.00", 16),
    ("Salary Position", "salary_position", None, 16),
    ("Compa Ratio", "compa_ratio", "0.00", 12),
    ("Increment %", "increment_pct", "0.00%", 12),
    ("Salary Correction %", "correction_pct", "0.00%", 18),
    ("Total Increase %", "total_increase_pct", "0.00%", 16),
    ("Recommended Salary", "recommended_salary", "0.00", 18),
    ("Increment Cost", "increment_amount", "0.00", 14),
    ("Salary Correction Cost", "correction_amount", "0.00", 20),
    ("Flags", "flags", None, 42),
]

COMPARISON_METRICS = [
    ("Employee ID", "employee_id", None),
    ("Employee Name", "employee_name", None),
    ("Department", "department", None),
    ("Grade", "grade", None),
    ("Current Salary", "current_salary", "0.00"),
    ("Performance Rating", "performance_rating", "0"),
    ("Years At Level", "years_at_level", "0.0"),
    ("Minimum Salary", "band_min", "0.00"),
    ("Median Salary", "band_median", "0.00"),
    ("Maximum Salary", "band_max", "0.00"),
    ("Salary Position", "salary_position", None),
    ("Compa Ratio", "compa_ratio", "0.00"),
    ("Increment %", "increment_pct", "0.00%"),
    ("Salary Correction %", "correction_pct", "0.00%"),
    ("Total Increase %", "total_increase_pct", "0.00%"),
    ("Recommended Salary", "recommended_salary", "0.00"),
]


def _write_header(sheet: Worksheet, titles: list[str], fill_hex: str = NAVY) -> None:
    fill = PatternFill("solid", fgColor=fill_hex)
    for index, title in enumerate(titles, start=1):
        cell = sheet.cell(1, index, title)
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(titles))}1"
    sheet.row_dimensions[1].height = 28


def _apply_cell(cell, number_format: str | None = None) -> None:
    cell.font = CELL_FONT
    cell.border = THIN
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if number_format:
        cell.number_format = number_format
        cell.font = MONO_FONT


def _row_fill(result: CompensationResult) -> PatternFill | None:
    if result.salary_position == "Below Minimum":
        return PatternFill("solid", fgColor=BELOW_MIN)
    if "High performer below market - priority correction" in result.flags:
        return PatternFill("solid", fgColor=PRIORITY)
    if result.salary_position == "Above Band":
        return PatternFill("solid", fgColor=ABOVE)
    return None


def _write_reports(sheet: Worksheet, reports: Iterable[CompensationResult], title_fill: str = NAVY) -> None:
    _write_header(sheet, [column[0] for column in REPORT_COLUMNS], title_fill)
    for row_index, result in enumerate(reports, start=2):
        data = result.to_dict()
        fill = _row_fill(result)
        for col_index, (_, field, number_format, _) in enumerate(REPORT_COLUMNS, start=1):
            value = data.get(field)
            if field == "flags":
                value = "; ".join(value or [])
            cell = sheet.cell(row_index, col_index, value)
            _apply_cell(cell, number_format)
            if fill:
                cell.fill = fill
    for index, (_, _, _, width) in enumerate(REPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"


def _write_table(sheet: Worksheet, rows: list[dict], columns: list[tuple[str, str, str | None, int]], fill_hex: str) -> None:
    _write_header(sheet, [column[0] for column in columns], fill_hex)
    for row_index, row in enumerate(rows, start=2):
        for col_index, (_, field, number_format, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, col_index, row.get(field))
            _apply_cell(cell, number_format)
    for index, (_, _, _, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_kv(sheet: Worksheet, items: list[tuple[str, object]]) -> None:
    _write_header(sheet, ["Metric", "Value"], COPPER)
    for row_index, (label, value) in enumerate(items, start=2):
        label_cell = sheet.cell(row_index, 1, label)
        value_cell = sheet.cell(row_index, 2, value)
        _apply_cell(label_cell)
        _apply_cell(value_cell)
        if isinstance(value, float) and "pct" in label.lower() or str(label).endswith("%"):
            value_cell.number_format = "0.00%"
            value_cell.font = MONO_FONT
        elif isinstance(value, float):
            value_cell.number_format = "0.00"
            value_cell.font = MONO_FONT
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 22


def build_compensation_workbook(
    reports: list[CompensationResult],
    budget: dict,
    issues: list[ValidationIssue] | None = None,
) -> bytes:
    workbook = Workbook()
    report_sheet = workbook.active
    report_sheet.title = "Compensation_Report"
    _write_reports(report_sheet, reports)

    summary = workbook.create_sheet("Budget_Summary")
    _write_kv(
        summary,
        [
            ("Scope", budget.get("scope")),
            ("Headcount", budget.get("headcount")),
            ("Total Current Payroll", budget.get("current_payroll")),
            ("Total Proposed Payroll", budget.get("recommended_payroll")),
            ("Increment Cost", budget.get("increment_cost")),
            ("Salary Correction Cost", budget.get("correction_cost")),
            ("Total Compensation Cost", budget.get("total_compensation_cost")),
            ("Average Increase %", budget.get("increase_pct")),
            ("Average Increment %", budget.get("avg_increment_pct")),
            ("Average Correction %", budget.get("avg_correction_pct")),
            ("Average Total Increase %", budget.get("avg_total_increase_pct")),
            ("Average Compa Ratio", budget.get("avg_compa_ratio")),
            ("Employees With Correction", budget.get("employees_with_correction")),
            ("Employees Below Minimum", budget.get("employees_below_minimum")),
            ("Employees Above Band", budget.get("employees_above_band")),
            ("Priority Cases", budget.get("priority_cases")),
        ],
    )

    group_columns = [
        ("Group", "label", None, 22),
        ("Headcount", "headcount", "0", 12),
        ("Current Payroll", "current_payroll", "0.00", 18),
        ("Proposed Payroll", "recommended_payroll", "0.00", 18),
        ("Increment Cost", "increment_cost", "0.00", 16),
        ("Salary Correction Cost", "correction_cost", "0.00", 20),
        ("Total Compensation Cost", "total_compensation_cost", "0.00", 22),
        ("Average Increase %", "increase_pct", "0.00%", 16),
        ("Avg Increment %", "avg_increment_pct", "0.00%", 16),
        ("Avg Correction %", "avg_correction_pct", "0.00%", 16),
        ("Avg Compa Ratio", "avg_compa_ratio", "0.00", 16),
    ]
    _write_table(workbook.create_sheet("Budget_By_Department"), budget.get("by_department") or [], group_columns, NAVY)
    _write_table(workbook.create_sheet("Budget_By_Grade"), budget.get("by_grade") or [], group_columns, NAVY)
    _write_table(workbook.create_sheet("Budget_By_Rating"), budget.get("by_rating") or [], group_columns, NAVY)
    _write_table(workbook.create_sheet("Budget_By_Position"), budget.get("by_position") or [], group_columns, COPPER)

    exceptions = workbook.create_sheet("Exceptions")
    flagged = [row for row in reports if row.flags]
    _write_reports(exceptions, flagged, COPPER)

    validation_sheet = workbook.create_sheet("Validation_Report")
    validation_rows = [
        {
            "sheet": item.sheet,
            "row": item.row,
            "employee_id": item.employee_id,
            "code": item.code,
            "severity": item.severity,
            "message": item.message,
        }
        for item in (issues or [])
    ]
    if not validation_rows:
        validation_rows = [
            {
                "sheet": "All",
                "row": None,
                "employee_id": None,
                "code": "OK",
                "severity": "info",
                "message": "No validation errors.",
            }
        ]
    _write_table(
        validation_sheet,
        validation_rows,
        [
            ("Sheet", "sheet", None, 24),
            ("Row", "row", "0", 10),
            ("Employee ID", "employee_id", None, 14),
            ("Code", "code", None, 20),
            ("Severity", "severity", None, 12),
            ("Message", "message", None, 64),
        ],
        COPPER,
    )

    return _save(workbook)


def build_validation_workbook(issues: list[ValidationIssue]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Validation_Report"
    columns = [
        ("Sheet", "sheet", None, 24),
        ("Row", "row", "0", 10),
        ("Employee ID", "employee_id", None, 14),
        ("Code", "code", None, 20),
        ("Severity", "severity", None, 12),
        ("Message", "message", None, 64),
    ]
    rows = [
        {
            "sheet": item.sheet,
            "row": item.row,
            "employee_id": item.employee_id,
            "code": item.code,
            "severity": item.severity,
            "message": item.message,
        }
        for item in issues
    ]
    if not rows:
        rows = [
            {
                "sheet": "All",
                "row": None,
                "employee_id": None,
                "code": "OK",
                "severity": "info",
                "message": "No validation errors.",
            }
        ]
    _write_table(sheet, rows, columns, COPPER)
    return _save(workbook)


def build_selected_workbook(reports: list[CompensationResult]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Selected_Employees"
    _write_reports(sheet, reports)
    return _save(workbook)


def build_comparison_workbook(reports: list[CompensationResult]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employee_Comparison"

    metrics = COMPARISON_METRICS

    titles = ["Metric"] + [row.employee_id for row in reports]
    _write_header(sheet, titles, COPPER)
    for row_index, (label, field, number_format) in enumerate(metrics, start=2):
        label_cell = sheet.cell(row_index, 1, label)
        _apply_cell(label_cell)
        label_cell.font = Font(name="Calibri", bold=True, color=INK, size=11)
        for col_index, result in enumerate(reports, start=2):
            data = result.to_dict()
            value = data.get(field)
            if field == "flags":
                value = "; ".join(value or [])
            cell = sheet.cell(row_index, col_index, value)
            _apply_cell(cell, number_format)
    sheet.column_dimensions["A"].width = 28
    for index in range(2, len(reports) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = 22
    sheet.page_setup.orientation = "landscape"
    return _save(workbook)


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
