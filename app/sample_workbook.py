from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="243044")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="D7CBB8"),
    right=Side(style="thin", color="D7CBB8"),
    top=Side(style="thin", color="D7CBB8"),
    bottom=Side(style="thin", color="D7CBB8"),
)

EMPLOYEE_HEADERS = [
    "Employee ID",
    "Employee Name",
    "Department",
    "Grade",
    "Designation",
    "Current Salary",
    "Performance Rating",
    "Date of Joining",
    "Current Grade Since",
    "Years At Level",
    "Total Years Experience",
    "Location",
    "Reporting Manager",
    "Reporting Partner",
    "Employment Status",
]

INCREMENT_HEADERS = ["Grade", "Performance Rating", "Increment Percentage"]
BAND_HEADERS = [
    "Department",
    "Grade",
    "Years At Level",
    "Minimum Salary",
    "Median Salary",
    "Maximum Salary",
]
CORRECTION_HEADERS = ["Minimum Compa Ratio", "Maximum Compa Ratio", "Salary Correction %"]

EMPLOYEES = [
    # Spec example: 10 LPA vs 14 LPA median, rating 4 → 10% + 8% = 18% → 11.8 LPA
    [1001, "Ananya Rao", "Technology", "A1", "Software Engineer", 10, 4, date(2023, 4, 10), date(2024, 4, 1), 2, 4.5, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1002, "Rahul Mehta", "Technology", "A1", "Software Engineer", 14, 4, date(2022, 6, 1), date(2024, 4, 1), 2, 5.0, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1003, "Priya Nair", "Technology", "A1", "Software Engineer", 9, 5, date(2023, 7, 15), date(2024, 4, 1), 2, 3.5, "Hyderabad", "Vikram Shah", "Neha Kapoor", "Active"],
    [1004, "Arjun Iyer", "Technology", "A1", "Software Engineer", 14, 3, date(2021, 1, 11), date(2024, 4, 1), 2, 6.0, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1005, "Meera Joshi", "Technology", "A1", "Senior Engineer", 16, 4, date(2020, 8, 3), date(2024, 4, 1), 2, 7.0, "Pune", "Vikram Shah", "Neha Kapoor", "Active"],
    [1006, "Sanjay Kulkarni", "Technology", "A1", "Senior Engineer", 21, 3, date(2018, 2, 19), date(2024, 4, 1), 2, 9.0, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1007, "Diya Banerjee", "Technology", "A1", "Software Engineer", 11.9, 4, date(2023, 1, 9), date(2024, 4, 1), 2, 4.0, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1008, "Karthik Reddy", "Technology", "A1", "Software Engineer", 13.3, 2, date(2022, 11, 21), date(2024, 4, 1), 2, 4.2, "Hyderabad", "Vikram Shah", "Neha Kapoor", "Active"],
    [1009, "Ishita Malhotra", "Technology", "A2", "Tech Lead", 18, 5, date(2019, 5, 6), date(2023, 4, 1), 3, 8.0, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1010, "Rohan Desai", "Technology", "A2", "Tech Lead", 22, 4, date(2018, 9, 12), date(2023, 4, 1), 3, 9.5, "Mumbai", "Vikram Shah", "Neha Kapoor", "Active"],
    [1011, "Sneha Kulkarni", "Technology", "B1", "Engineering Manager", 28, 4, date(2016, 3, 1), date(2022, 4, 1), 4, 12.0, "Bengaluru", "Neha Kapoor", "Amit Bansal", "Active"],
    [1012, "Farhan Qureshi", "Technology", "B1", "Engineering Manager", 24, 5, date(2017, 7, 18), date(2022, 4, 1), 4, 11.0, "Bengaluru", "Neha Kapoor", "Amit Bansal", "Active"],
    [1013, "Leela Krishnan", "Finance", "A1", "Financial Analyst", 8.5, 4, date(2023, 2, 1), date(2024, 4, 1), 2, 3.8, "Mumbai", "Kavita Rao", "Amit Bansal", "Active"],
    [1014, "Nikhil Sharma", "Finance", "A1", "Financial Analyst", 12, 3, date(2022, 4, 4), date(2024, 4, 1), 2, 4.5, "Mumbai", "Kavita Rao", "Amit Bansal", "Active"],
    [1015, "Aisha Khan", "Finance", "A2", "Finance Manager", 16.5, 5, date(2019, 10, 14), date(2023, 4, 1), 3, 8.0, "Mumbai", "Kavita Rao", "Amit Bansal", "Active"],
    [1016, "Vivek Pillai", "Human Resources", "A1", "HR Executive", 7.5, 3, date(2023, 5, 2), date(2024, 4, 1), 2, 3.0, "Bengaluru", "Sonal Gupta", "Amit Bansal", "Active"],
    [1017, "Tara Menon", "Human Resources", "A2", "HR Business Partner", 14, 4, date(2020, 1, 20), date(2023, 4, 1), 3, 7.0, "Bengaluru", "Sonal Gupta", "Amit Bansal", "Active"],
    [1018, "Manish Patel", "Sales", "A1", "Account Executive", 9.5, 5, date(2023, 8, 7), date(2024, 4, 1), 2, 3.2, "Delhi", "Ritu Agarwal", "Amit Bansal", "Active"],
    [1019, "Pooja Sinha", "Sales", "A1", "Account Executive", 13, 2, date(2021, 12, 13), date(2024, 4, 1), 2, 5.5, "Delhi", "Ritu Agarwal", "Amit Bansal", "Active"],
    [1020, "Aditya Bose", "Sales", "A2", "Sales Manager", 17, 4, date(2018, 6, 25), date(2023, 4, 1), 3, 9.0, "Mumbai", "Ritu Agarwal", "Amit Bansal", "Active"],
    [1021, "Hannah D'Souza", "Operations", "A1", "Ops Analyst", 8, 4, date(2024, 1, 8), date(2024, 4, 1), 2, 2.8, "Pune", "Deepak Jain", "Amit Bansal", "Active"],
    [1022, "Gaurav Jain", "Operations", "A2", "Ops Lead", 13.5, 3, date(2019, 9, 30), date(2023, 4, 1), 3, 7.5, "Pune", "Deepak Jain", "Amit Bansal", "Active"],
    [1023, "Nandini Rao", "Technology", "A1", "Software Engineer", 10.5, 1, date(2023, 3, 14), date(2024, 4, 1), 2, 4.0, "Bengaluru", "Vikram Shah", "Neha Kapoor", "Active"],
    [1024, "Omar Sheikh", "Technology", "C1", "Principal Engineer", 36, 5, date(2014, 4, 2), date(2021, 4, 1), 5, 14.0, "Bengaluru", "Neha Kapoor", "Amit Bansal", "Active"],
    [1025, "Kavya Sharma", "Technology", "A1", "Software Engineer", 10, 4, date(2022, 2, 28), date(2024, 4, 1), 2, 4.1, "Hyderabad", "Vikram Shah", "Neha Kapoor", "Notice Period"],
    [1026, "Ritika Sen", "Finance", "B1", "Controller", 26, 4, date(2015, 11, 9), date(2022, 4, 1), 4, 13.0, "Mumbai", "Kavita Rao", "Amit Bansal", "Active"],
    [1027, "Harish Venkatesh", "Technology", "A2", "Tech Lead", 15, 3, date(2020, 7, 1), date(2024, 4, 1), 2, 6.5, "Chennai", "Vikram Shah", "Neha Kapoor", "Active"],
    [1028, "Sana Iqbal", "Sales", "B1", "Regional Head", 30, 5, date(2014, 8, 18), date(2021, 4, 1), 5, 14.5, "Delhi", "Ritu Agarwal", "Amit Bansal", "Inactive"],
]

INCREMENT_GRID = []
GRADE_INCREMENTS = {
    "A1": {5: 0.15, 4: 0.10, 3: 0.07, 2: 0.03, 1: 0.00},
    "A2": {5: 0.14, 4: 0.10, 3: 0.06, 2: 0.03, 1: 0.00},
    "B1": {5: 0.12, 4: 0.09, 3: 0.06, 2: 0.02, 1: 0.00},
    "B2": {5: 0.11, 4: 0.08, 3: 0.05, 2: 0.02, 1: 0.00},
    "C1": {5: 0.10, 4: 0.08, 3: 0.05, 2: 0.02, 1: 0.00},
}
for grade, mapping in GRADE_INCREMENTS.items():
    for rating, pct in mapping.items():
        INCREMENT_GRID.append([grade, rating, f"{int(pct * 100)}%"])

BANDS = [
    ["Technology", "A1", "0 Years", 9, 12, 16],
    ["Technology", "A1", "1 Years", 10, 13, 18],
    ["Technology", "A1", "2 Years", 11, 14, 20],
    ["Technology", "A1", "3 Years", 12, 15.5, 22],
    ["Technology", "A2", "2 Years", 16, 20, 28],
    ["Technology", "A2", "3 Years", 17, 21, 30],
    ["Technology", "B1", "4 Years", 24, 30, 40],
    ["Technology", "C1", "5 Years", 32, 40, 55],
    ["Finance", "A1", "2 Years", 8, 11, 16],
    ["Finance", "A2", "3 Years", 14, 18, 26],
    ["Finance", "B1", "4 Years", 22, 28, 38],
    ["Human Resources", "A1", "2 Years", 7, 9.5, 14],
    ["Human Resources", "A2", "3 Years", 12, 15, 22],
    ["Sales", "A1", "2 Years", 8, 12, 18],
    ["Sales", "A2", "3 Years", 14, 18, 26],
    ["Sales", "B1", "5 Years", 24, 32, 45],
    ["Operations", "A1", "2 Years", 7.5, 10, 15],
    ["Operations", "A2", "3 Years", 12, 15.5, 22],
]

CORRECTIONS = [
    [0.00, 0.80, "8%"],
    [0.80, 0.90, "5%"],
    [0.90, 1.00, "2%"],
    [1.00, 999, "0%"],
]


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, col_index, value)
            cell.border = THIN
            if isinstance(value, date):
                cell.number_format = "DD-MMM-YYYY"
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 4)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def build_sample_workbook() -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _write_sheet(workbook, "Employee_Master", EMPLOYEE_HEADERS, EMPLOYEES)
    _write_sheet(workbook, "Increment_Grid", INCREMENT_HEADERS, INCREMENT_GRID)
    _write_sheet(workbook, "Salary_Band", BAND_HEADERS, BANDS)
    _write_sheet(workbook, "Salary_Correction_Rules", CORRECTION_HEADERS, CORRECTIONS)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
